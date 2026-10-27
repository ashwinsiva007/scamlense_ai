import io
import pytest
from PIL import Image
import docx
from docx.opc.constants import RELATIONSHIP_TYPE as RT
import email.message
import fitz

from app.services.file_utils import (
    process_uploaded_file,
    detect_file_type_magic,
    FileValidationError,
    validate_file_integrity_and_type,
)
from app.config import settings


def create_text_pdf() -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), "Congratulations! Selected for Internship. Pay Rs 1500 to secure seat.")
    rect = fitz.Rect(50, 72, 350, 90)
    page.insert_link({"kind": fitz.LINK_URI, "from": rect, "uri": "https://suspicious-job-portal.com/verify"})
    buf = doc.tobytes()
    doc.close()
    return buf


def create_scanned_pdf() -> bytes:
    # A PDF containing purely a raster image and zero selectable text
    img = Image.new("RGB", (400, 300), color=(200, 200, 200))
    img_io = io.BytesIO()
    img.save(img_io, format="JPEG")
    doc = fitz.open()
    page = doc.new_page(width=400, height=300)
    page.insert_image(fitz.Rect(0, 0, 400, 300), stream=img_io.getvalue())
    buf = doc.tobytes()
    doc.close()
    return buf


def create_docx_with_link() -> bytes:
    doc = docx.Document()
    doc.add_paragraph("Please confirm your admission by clicking the link below:")
    doc.part.relate_to("https://admissions-fee-portal.net/pay", RT.HYPERLINK, is_external=True)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Security Fee"
    table.rows[0].cells[1].text = "INR 2,500"
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def create_eml_with_link() -> bytes:
    msg = email.message.EmailMessage()
    msg["From"] = "security@alert-sbi-portal.com"
    msg["To"] = "victim@example.com"
    msg["Subject"] = "URGENT: Your bank account is blocked"
    msg.set_content(
        "Dear Customer,\n\nYour account has been restricted. Visit https://sbi-unblock-secure.com immediately."
    )
    return msg.as_bytes()


def create_txt_sample() -> bytes:
    return "Electricity power will be disconnected tonight at 9:30 PM. Call electricity officer at 9876543210.".encode("utf-8")


def create_png_sample() -> bytes:
    img = Image.new("RGB", (300, 200), color=(100, 150, 200))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def test_text_pdf_extraction():
    pdf_bytes = create_text_pdf()
    assert detect_file_type_magic(pdf_bytes) == "PDF"
    result = process_uploaded_file(pdf_bytes, "offer.pdf")
    assert result.detected_format == "PDF"
    assert "Internship" in result.extracted_text
    assert any("suspicious-job-portal.com" in link for link in result.hyperlinks)


def test_scanned_pdf_rendering():
    scanned_bytes = create_scanned_pdf()
    assert detect_file_type_magic(scanned_bytes) == "PDF"
    result = process_uploaded_file(scanned_bytes, "scan.pdf")
    assert result.detected_format == "PDF"
    # Scanned PDF should render at least 1 image
    assert len(result.images) >= 1
    assert result.images[0][0].startswith(b"\xff\xd8\xff")  # valid JPEG


def test_docx_extraction_with_link():
    docx_bytes = create_docx_with_link()
    assert detect_file_type_magic(docx_bytes) == "DOCX"
    result = process_uploaded_file(docx_bytes, "notice.docx")
    assert result.detected_format == "DOCX"
    assert "Security Fee" in result.extracted_text
    assert any("admissions-fee-portal.net" in link for link in result.hyperlinks)


def test_eml_extraction():
    eml_bytes = create_eml_with_link()
    assert detect_file_type_magic(eml_bytes) == "EML"
    result = process_uploaded_file(eml_bytes, "phishing.eml")
    assert result.detected_format == "EML"
    assert "security@alert-sbi-portal.com" in result.sender
    assert "URGENT" in result.subject
    assert any("sbi-unblock-secure.com" in link for link in result.hyperlinks)


def test_txt_extraction():
    txt_bytes = create_txt_sample()
    assert detect_file_type_magic(txt_bytes) == "TXT"
    result = process_uploaded_file(txt_bytes, "alert.txt")
    assert result.detected_format == "TXT"
    assert "Electricity power" in result.extracted_text


def test_png_image_processing():
    png_bytes = create_png_sample()
    assert detect_file_type_magic(png_bytes) == "PNG"
    result = process_uploaded_file(png_bytes, "sample.png")
    assert result.detected_format == "PNG"
    assert len(result.images) == 1
    assert result.images[0][0].startswith(b"\xff\xd8\xff")  # converted to JPEG


def test_mislabelled_file_rejected():
    # Plain text file disguised as .png
    fake_png = b"This is plain text claiming to be a PNG image."
    with pytest.raises(FileValidationError) as exc_info:
        validate_file_integrity_and_type(fake_png, "fake.png")
    assert "does not match its claimed extension" in str(exc_info.value)


def test_oversized_file_rejected():
    # Exceeds MAX_UPLOAD_MB
    oversized = b"0" * ((settings.MAX_UPLOAD_MB * 1024 * 1024) + 1024)
    with pytest.raises(FileValidationError) as exc_info:
        validate_file_integrity_and_type(oversized, "huge.txt")
    assert "exceeds maximum allowed limit" in str(exc_info.value)


def test_corrupt_pdf_rejected():
    corrupt_pdf = b"%PDF-1.4\nInvalid content" + b"\x00\xFF" * 100
    with pytest.raises(FileValidationError) as exc_info:
        process_uploaded_file(corrupt_pdf, "broken.pdf")
    assert "Could not read document" in str(exc_info.value) or "Corrupt" in str(exc_info.value)
